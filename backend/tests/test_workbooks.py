import io
import json
import zipfile
from pathlib import Path
from uuid import uuid4
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from ledger.workbooks import file_hash, process_job
from conftest import mutation


def fixture_bytes():
    wb=Workbook(); ws=wb.active; ws.title='Expenses'
    ws.append(['Date','Category','Amount','Description','Vendor','Method','Kind','Double'])
    ws.append(['2026-10-01','Cement',12500.25,'Cement delivery','Supplier','upi','expense','=C2*2'])
    ws['C2'].fill=PatternFill('solid',fgColor='00CCFF')
    summary=wb.create_sheet('Summary'); summary['A1']='=SUM(Expenses!C:C)'
    result=io.BytesIO(); wb.save(result); return result.getvalue()


def link(env,data=None,columns=None):
    headers=env['tokens']['admin']
    uploaded=env['client'].post('/workbooks/inspect',headers=headers,files={'file':('costs.xlsx',data or fixture_bytes())})
    assert uploaded.status_code==200, uploaded.text
    mapped={'upload_id':uploaded.json()['upload_id'],'sheet':'Expenses','header_row':1,
            'columns': columns or {'expense_date':'A','category':'B','amount_minor':'C','description':'D','vendor':'E','payment_method':'F','kind':'G'}}
    return env['client'].post('/projects/'+env['project']+'/workbook/link',headers=headers,json=mapped)


def path(env):
    return env['app'].state.database.settings.data_dir/'workbooks'/f'{env["project"]}.xlsx'


def test_import_add_edit_refund_void_preserve_formulas_styles(env):
    assert link(env).status_code==200
    workbook=load_workbook(path(env)); ws=workbook['Expenses']
    assert ws['H2'].value=='=C2*2' and workbook['Summary']['A1'].value=='=SUM(Expenses!C:C)'
    assert ws['C2'].fill.fgColor.rgb=='0000CCFF'
    body=mutation(env)
    response=env['client'].post('/sync/mutate',headers=env['tokens']['editor'],json=body)
    assert response.status_code==200
    workbook=load_workbook(path(env)); ws=workbook['Expenses']
    assert ws['C3'].value==3000 and ws['H3'].value=='=C3*2'
    assert ws['C3'].fill.fgColor.rgb=='0000CCFF'
    update={**body,'operation_id':str(uuid4()),'action':'update','base_version':1,'expense':{**body['expense'],'amount_minor':325001,'description':'=HYPERLINK("bad")'}}
    assert env['client'].post('/sync/mutate',headers=env['tokens']['editor'],json=update).status_code==200
    ws=load_workbook(path(env))['Expenses']
    assert ws['C3'].value==3250.01 and ws['D3'].data_type=='s'
    reverse={**update,'operation_id':str(uuid4()),'action':'void','base_version':2,'reason':'Wrong transport bill'}
    assert env['client'].post('/sync/mutate',headers=env['tokens']['editor'],json=reverse).status_code==200
    ws=load_workbook(path(env))['Expenses']
    assert ws['C3'].value==0 and ws['K3'].value=='void'
    assert list(path(env).parent.glob('*.original.xlsx'))
    assert list(path(env).parent.glob('*.v*.xlsx'))


def test_external_changes_block_without_losing_database_expense(env):
    assert link(env).status_code==200
    workbook=load_workbook(path(env)); workbook['Summary']['B1']='External change'; workbook.save(path(env))
    original=file_hash(path(env))
    assert env['client'].post('/sync/mutate',headers=env['tokens']['editor'],json=mutation(env)).status_code==200
    assert file_hash(path(env))==original
    pull=env['client'].get('/sync/pull',headers=env['tokens']['editor']).json()
    assert len(pull['snapshot'])==2
    assert pull['workbooks'][0]['status']=='blocked'
    assert 'External' in pull['workbooks'][0]['error']


def test_interrupted_write_retry(env):
    assert link(env).status_code==200
    original=file_hash(path(env))
    real_replace=__import__('os').replace
    def failing_replace(source,target):
        if str(target)==str(path(env)):
            raise OSError('simulated disk write failure')
        return real_replace(source,target)
    with patch('ledger.workbooks.os.replace',side_effect=failing_replace):
        assert env['client'].post('/sync/mutate',headers=env['tokens']['editor'],json=mutation(env)).status_code==200
    assert file_hash(path(env))==original
    assert process_job(env['app'].state.database,env['project']) is True
    assert load_workbook(path(env))['Expenses']['C3'].value==3000


def test_office_lock_retries(env):
    assert link(env).status_code==200
    lock=path(env).parent/('~$'+path(env).name); lock.touch()
    original=file_hash(path(env))
    assert env['client'].post('/sync/mutate',headers=env['tokens']['editor'],json=mutation(env)).status_code==200
    assert file_hash(path(env))==original
    lock.unlink()
    assert process_job(env['app'].state.database,env['project'])


def test_unsupported_and_corrupt_workbooks_rejected(env):
    response=env['client'].post('/workbooks/inspect',headers=env['tokens']['admin'],files={'file':('bad.xlsx',b'notzip')})
    assert response.status_code==422
    data=io.BytesIO(fixture_bytes())
    with zipfile.ZipFile(data,'a') as package:
        package.writestr('xl/drawings/drawing1.xml','<drawing/>')
    response=env['client'].post('/workbooks/inspect',headers=env['tokens']['admin'],files={'file':('advanced.xlsx',data.getvalue())})
    assert response.status_code==422 and 'Unsupported' in response.text


def test_invalid_mapping_rolls_back_import(env):
    response=link(env,columns={'expense_date':'A','category':'B','amount_minor':'B','description':'D'})
    assert response.status_code==422
    assert env['client'].get('/sync/pull',headers=env['tokens']['admin']).json()['snapshot']==[]
    assert not path(env).exists()


def test_external_reconciliation_preserves_reviewed_unmapped_edits(env):
    assert link(env).status_code==200
    workbook=load_workbook(path(env)); workbook['Summary']['B1']='Reviewed note'; workbook.save(path(env))
    reviewed=file_hash(path(env))
    url='/projects/'+env['project']+'/workbook/reconcile'
    result=env['client'].post(url,headers=env['tokens']['admin'],json={'reviewed_sha256':reviewed,'confirmed':True})
    assert result.status_code==200, result.text
    assert load_workbook(path(env))['Summary']['B1'].value=='Reviewed note'


def test_external_financial_difference_never_adopted(env):
    assert link(env).status_code==200
    workbook=load_workbook(path(env)); workbook['Expenses']['C2']=1; workbook.save(path(env))
    reviewed=file_hash(path(env))
    result=env['client'].post('/projects/'+env['project']+'/workbook/reconcile',headers=env['tokens']['admin'],json={'reviewed_sha256':reviewed,'confirmed':True})
    assert result.status_code==409
    assert result.json()['detail']['code']=='external_financial_differences'
    assert file_hash(path(env))==reviewed


def test_replay_after_replacement_before_database_hash_commit(env):
    assert link(env).status_code==200
    import ledger.workbooks as workbooks
    real_fsync=workbooks.fsync_dir
    calls=0
    def crash_after_replace(directory):
        nonlocal calls
        calls+=1
        if calls==2:
            raise OSError('simulated interruption after workbook replace before hash commit')
        return real_fsync(directory)
    with patch('ledger.workbooks.fsync_dir',side_effect=crash_after_replace):
        assert env['client'].post('/sync/mutate',headers=env['tokens']['editor'],json=mutation(env)).status_code==200
    assert load_workbook(path(env))['Expenses']['C3'].value==3000
    assert process_job(env['app'].state.database,env['project']) is True
    ws=load_workbook(path(env))['Expenses']
    assert ws.max_row==3 and ws['C3'].value==3000
    pull=env['client'].get('/sync/pull',headers=env['tokens']['editor']).json()
    assert pull['workbooks'][0]['status']=='synced'
