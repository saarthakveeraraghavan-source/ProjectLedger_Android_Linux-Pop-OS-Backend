"""Run native acceptance against a temporary private backend; stop the synthetic backend afterward. The operator owns the separate temporary route."""
import json
import os
import subprocess
import tempfile
import time
from pathlib import Path
from uuid import uuid4

import httpx
from openpyxl import load_workbook
from ledger.auth import create_user
from ledger.backup import create_backup, restore_backup
from ledger.config import Settings
from ledger.db import Database
from ledger.workbooks import drain_jobs

root=Path(__file__).resolve().parents[1]
adb=str(root/'.tools/sdk/platform-tools/adb')
serial='AGN2DK23IB037889'
base='https://pop-os.tail4722e9.ts.net:8443'
password=str(uuid4())+str(uuid4())
current=json.loads(subprocess.check_output(['tailscale','serve','status','--json']))
handler = current.get('Web',{}).get('pop-os.tail4722e9.ts.net:8443',{}).get('Handlers',{}).get('/',{})
if handler.get('Proxy') != 'http://127.0.0.1:18766':
    raise RuntimeError('Existing private route differs; do not alter tailnet configuration')
import socket
with socket.socket() as check:
    check.bind(('127.0.0.1',18766))
# Use the already configured route only when its local port is vacant. No Serve mutation.

with tempfile.TemporaryDirectory(prefix='phone-acceptance-',dir=root/'.tools') as temporary:
    settings=Settings(data_dir=Path(temporary)/'data')
    database=Database(settings)
    with database.transaction() as db: create_user(db,'acceptance-admin',password,True)
    environment={**os.environ,'LEDGER_DATA_DIR':str(settings.data_dir),'PYTHONPATH':str(root/'backend')}
    log=open(root/'docs/phone-acceptance-server.log','w')
    process=subprocess.Popen([str(root.parent/'.venv/bin/python'),'-m','uvicorn','ledger.api:create_app','--factory','--host','127.0.0.1','--port','18766','--no-access-log'],env=environment,stdout=log,stderr=log)
    configured=False
    try:
        with httpx.Client(timeout=10,trust_env=False) as client:
            for _ in range(100):
                try:
                    if client.get('http://127.0.0.1:18766/health').status_code==200: break
                except httpx.HTTPError: pass
                time.sleep(.1)
            else: raise RuntimeError('Temporary backend did not become ready')
        for apk in [root/'android/app/build/outputs/apk/debug/app-debug.apk',root/'android/app/build/outputs/apk/androidTest/debug/app-debug-androidTest.apk']:
            subprocess.run([adb,'-s',serial,'install','-r',str(apk)],check=True)
        args=[adb,'-s',serial,'shell','am','instrument','-w','-e','class','com.projectledger.ConnectedAcceptanceTest',
              '-e','ledgerBase',base,'-e','ledgerUsername','acceptance-admin','-e','ledgerPassword',password,
              'com.projectledger.test/androidx.test.runner.AndroidJUnitRunner']
        try:
            result=subprocess.run(args,capture_output=True,text=True,timeout=600)
        except subprocess.TimeoutExpired:
            raise RuntimeError('Native acceptance timed out; temporary test backend will be stopped') from None
        output=result.stdout+'\n'+result.stderr
        (root/'docs/android-phone-acceptance.txt').write_text(output)
        if 'OK (1 test)' not in output: raise RuntimeError('Native acceptance failed; inspect docs/android-phone-acceptance.txt')
        drain_jobs(database)
        with database.connect() as db:
            project=db.execute('SELECT id FROM projects').fetchone()[0]
            linked=db.execute('SELECT * FROM workbooks WHERE project_id=?',(project,)).fetchone()
            assert linked['status']=='synced',dict(linked)
            wb=load_workbook(linked['path']);ws=wb['Expenses']
            amounts=[ws.cell(r,3).value for r in range(2,ws.max_row+1)]
            from decimal import Decimal
            net=sum(Decimal(str(a)) for a in amounts)
            assert net==Decimal('17400.01'),amounts
            assert wb['Summary']['A1'].value=='=SUM(Expenses!C:C)'
            assert db.execute('SELECT COUNT(*) FROM expenses').fetchone()[0]==4
        backup=create_backup(database,Path(temporary)/'backup')
        restored=restore_backup(backup,Path(temporary)/'restored')
        restored_db=Database(Settings(data_dir=restored))
        with restored_db.connect() as db:
            assert db.execute('SELECT COUNT(*) FROM expenses').fetchone()[0]==4
            assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        report={'device':'LAVA LXX504 Android API 35','transport':'operator-configured Tailscale private HTTPS 8443','native_acceptance':'passed',
                'workbook_status':'synced','workbook_net_inr':'17400.01','formula_preserved':True,'committed_records':4,
                'backup_restore_integrity':'passed','data':'synthetic only','test_script_does_not_modify_tailscale':True}
        (root/'docs/phone-acceptance-verification.json').write_text(json.dumps(report,indent=2))
        print(json.dumps(report),flush=True)
    finally:
        if configured: subprocess.run(['tailscale','serve','--https=8443','off'],check=True)
        process.terminate()
        try:process.wait(timeout=10)
        except subprocess.TimeoutExpired:process.kill();process.wait()
        log.close()
